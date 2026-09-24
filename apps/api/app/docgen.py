"""Document model plumbing + export: upgrade legacy docs to DocV2, and render a
DocV2 as Markdown, PDF (fpdf2) or DOCX (python-docx).

Why fonts are bundled: fpdf2's core fonts are Latin-1 only, so any non-Latin
title used to come out as "?". DejaVu Sans (regular + bold) ships under
app/fonts/ so Unicode works identically in `make dev` on macOS and in the slim
Debian image, which has no system fonts. If the files ever go missing we fall
back to Helvetica + `_latin()` rather than failing the export.

Inline markup: doc text carries only `**bold**` (UI labels) and `` `code` ``
(typed values). Each renderer maps exactly those two; nothing else is parsed.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from typing import Any

from app.schemas import DocV2
from app.storage import store

log = logging.getLogger("refract.docgen")

_FONT_DIR = Path(__file__).parent / "fonts"
_FONT_CANDIDATES = [
    (_FONT_DIR / "DejaVuSans.ttf", _FONT_DIR / "DejaVuSans-Bold.ttf"),
    (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ),
]

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


# ------------------------------------------------------------- model ----


def upgrade_doc(doc_json: dict[str, Any] | None, graph_json: dict[str, Any] | None = None) -> dict[str, Any]:
    """Return a validated DocV2 dict. v2 passes through; the legacy v1 shape
    {title, summary, graph_version, steps:[{n,title,body,screenshot}]} is mapped.
    v1 never stored graph step ids, so `source.graph_step_id` is filled by index
    only when the graph has the same number of steps (v1 was 1:1 with it).
    Idempotent."""
    doc_json = doc_json or {}
    if doc_json.get("version") == 2:
        return DocV2.model_validate(doc_json).model_dump()

    gsteps = (graph_json or {}).get("steps") or []
    v1_steps = doc_json.get("steps") or []
    aligned = len(gsteps) == len(v1_steps) and len(gsteps) > 0
    steps = []
    for i, s in enumerate(v1_steps, start=1):
        key = s.get("screenshot")
        g = gsteps[i - 1] if aligned else {}
        steps.append(
            {
                "id": f"legacy_{s.get('n', i)}",
                "title": s.get("title") or f"Step {i}",
                "body": s.get("body") or "",
                "tip": None,
                "snapshot": {"key": key, "raw_key": key, "t": g.get("t_start")} if key else None,
                "source": {
                    "graph_step_id": g.get("id") if aligned else None,
                    "t_start": g.get("t_start"),
                    "t_end": g.get("t_end"),
                },
            }
        )
    return DocV2.model_validate(
        {
            "version": 2,
            "title": doc_json.get("title") or "Workflow",
            "overview": doc_json.get("summary") or "",
            "steps": steps,
            "meta": {"writer": "legacy", "graph_version": doc_json.get("graph_version")},
        }
    ).model_dump()


def build_document_v2(
    graph_json: dict[str, Any], *, project_title: str | None = None, viewport: dict[str, Any] | None = None
) -> dict[str, Any]:
    """Mechanical v2 doc straight from a graph (no LLM, no ffmpeg) with the
    graph's own keyframes as snapshots. The public share route uses this when a
    project has no generated document."""
    from app.docwriter import attach_graph_screenshots, write_document_fallback

    steps = graph_json.get("steps", [])
    doc = write_document_fallback(
        project_title=project_title or graph_json.get("title") or "Workflow",
        transcript_text=" ".join((s.get("narration") or "").strip() for s in steps).strip(),
        steps=steps,
        graph_version=graph_json.get("version"),
    )
    return attach_graph_screenshots(doc, graph_json, viewport)


# ---------------------------------------------------------- markdown ----


def render_markdown(doc: dict[str, Any], media_base: str = "") -> str:
    doc = upgrade_doc(doc)
    lines = [f"# {doc['title']}", ""]
    if doc.get("overview"):
        lines += [doc["overview"], ""]
    if doc.get("prerequisites"):
        lines += ["## Prerequisites", ""] + [f"- {p}" for p in doc["prerequisites"]] + [""]
    for n, s in enumerate(doc["steps"], start=1):
        lines += [f"## {n}. {s['title']}", ""]
        if s.get("body"):
            lines += [s["body"], ""]
        if s.get("tip"):
            lines += [f"> **Tip:** {s['tip']}", ""]
        key = (s.get("snapshot") or {}).get("key")
        if key:
            url = f"{media_base}/media/{key}" if media_base else key
            lines += [f"![Step {n}]({url})", ""]
    if doc.get("tips"):
        lines += ["## Tips & troubleshooting", ""] + [f"- {t}" for t in doc["tips"]] + [""]
    return "\n".join(lines).strip() + "\n"


# --------------------------------------------------------------- pdf ----


def _latin(text: str) -> str:
    """fpdf core fonts are latin-1; map common unicode then drop the rest."""
    text = (
        (text or "")
        .replace("→", "->")
        .replace("—", "-")
        .replace("’", "'")
        .replace("“", '"')
        .replace("”", '"')
    )
    return text.encode("latin-1", "replace").decode("latin-1")


def resolve_pdf_font() -> tuple[Path, Path] | None:
    for regular, bold in _FONT_CANDIDATES:
        if regular.exists() and bold.exists():
            return regular, bold
    return None


def _pdf_inline(text: str) -> str:
    """Our markup -> fpdf2's: **bold** is native; code spans lose their ticks."""
    return re.sub(r"`([^`]+)`", r"\1", text or "")


def render_pdf(doc: dict[str, Any]) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    doc = upgrade_doc(doc)
    pdf = FPDF(format="A4")
    pdf.set_margins(16, 16, 16)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()

    fonts = resolve_pdf_font()
    if fonts:
        pdf.add_font("Doc", "", str(fonts[0]))
        pdf.add_font("Doc", "B", str(fonts[1]))
        pdf.add_font("Doc", "I", str(fonts[0]))  # no oblique bundled: same face
        family, txt = "Doc", (lambda s: s)
    else:  # pragma: no cover - fonts are shipped with the package
        family, txt = "Helvetica", _latin

    def para(size: float, style: str, text: str, h: float = 6, color=(20, 20, 20), markdown=True) -> None:
        pdf.set_font(family, style, size)
        pdf.set_text_color(*color)
        # new_x=LMARGIN keeps the cursor at the left margin (fpdf2 defaults to RIGHT).
        pdf.multi_cell(0, h, txt(_pdf_inline(text)), new_x=XPos.LMARGIN, new_y=YPos.NEXT, markdown=markdown)

    para(20, "B", doc["title"], h=10, markdown=False)
    if doc.get("overview"):
        para(11, "", doc["overview"], color=(90, 90, 90))
    if doc.get("prerequisites"):
        pdf.ln(3)
        para(14, "B", "Prerequisites", h=8, markdown=False)
        for p in doc["prerequisites"]:
            para(11, "", f"• {p}")
    pdf.ln(3)
    usable_w = pdf.epw
    for n, s in enumerate(doc["steps"], start=1):
        para(13, "B", f"{n}. {s['title']}", h=8, markdown=False)
        if s.get("body"):
            para(11, "", s["body"])
        if s.get("tip"):
            para(10, "I", f"Tip: {s['tip']}", color=(30, 110, 110))
        key = (s.get("snapshot") or {}).get("key")
        if key:
            try:
                path = store.fetch(key)
                if path is not None:
                    pdf.ln(1)
                    pdf.image(str(path), w=min(usable_w, 150))
            except Exception as e:  # a bad image must not kill the export
                log.warning("pdf: skipping snapshot %s (%s)", key, e)
        pdf.ln(5)
    if doc.get("tips"):
        para(14, "B", "Tips & troubleshooting", h=8, markdown=False)
        for t in doc["tips"]:
            para(11, "", f"• {t}")
    return bytes(pdf.output())


# -------------------------------------------------------------- docx ----

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _add_inline(paragraph, text: str, *, italic: bool = False) -> None:
    for part in _INLINE.split(text or ""):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(part)
        if italic:
            run.italic = True


def render_docx(doc: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = upgrade_doc(doc)
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    usable_in = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    d.add_heading(doc["title"], level=0)
    if doc.get("overview"):
        _add_inline(d.add_paragraph(), doc["overview"])
    if doc.get("prerequisites"):
        d.add_heading("Prerequisites", level=1)
        for item in doc["prerequisites"]:
            _add_inline(d.add_paragraph(style="List Bullet"), item)

    for n, s in enumerate(doc["steps"], start=1):
        d.add_heading(f"{n}. {s['title']}", level=2)
        if s.get("body"):
            _add_inline(d.add_paragraph(), s["body"])
        if s.get("tip"):
            tp = d.add_paragraph()
            tp.paragraph_format.left_indent = Inches(0.3)
            lead = tp.add_run("Tip: ")
            lead.bold = True
            lead.italic = True
            _add_inline(tp, s["tip"], italic=True)
        key = (s.get("snapshot") or {}).get("key")
        if key:
            try:
                path = store.fetch(key)
                if path is not None:
                    from PIL import Image  # via fpdf2's dependency

                    with Image.open(path) as im:
                        w_px = im.size[0]
                    width_in = min(usable_in, w_px / 96.0)  # never wider than the text column
                    d.add_picture(str(path), width=Inches(width_in))  # height follows: aspect kept
                    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception as e:
                log.warning("docx: skipping snapshot %s (%s)", key, e)

    if doc.get("tips"):
        d.add_heading("Tips & troubleshooting", level=1)
        for t in doc["tips"]:
            _add_inline(d.add_paragraph(style="List Bullet"), t)

    footer = sec.footer.paragraphs[0]
    footer.text = "Generated with RevEase"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def slugify(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", (text or "document").lower()).strip("-") or "document"
